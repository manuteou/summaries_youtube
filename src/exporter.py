import os
import io
import base64
from datetime import datetime
import html
import markdown
from xhtml2pdf import pisa
from markdownify import markdownify as md
from utils import slugify, clean_markdown_text
from abc import ABC, abstractmethod

# --- Helper Functions (Shared Logic) ---

def load_logo_base64():
    """Charge le logo en base64 pour l'intégration dans les PDF."""
    logo_path = os.path.join(os.path.dirname(__file__), "images", "synthetIA_logo.png")
    if os.path.exists(logo_path):
        with open(logo_path, "rb") as f:
            return base64.b64encode(f.read()).decode("utf-8")
    return None

def load_css():
    css_path = os.path.join(os.path.dirname(__file__), "assets", "style.css")
    if os.path.exists(css_path):
        with open(css_path, "r", encoding="utf-8") as f:
            return f.read()
    return ""

def is_html(text: str) -> bool:
    """Detects if text is likely HTML."""
    if not text:
        return False
    text = text.strip()
    return (
        text.startswith("<") or 
        "<p>" in text or 
        "<div>" in text or 
        "<h1>" in text or 
        "<span" in text
    )

def to_html(content: str) -> str:
    """Converts content to HTML."""
    if not content:
        return ""
    if is_html(content):
        return content
    else:
        return markdown.markdown(clean_markdown_text(content), extensions=['extra', 'codehilite'])

def to_markdown(content: str) -> str:
    """Converts content to Markdown."""
    if not content:
        return ""
    if is_html(content):
        return md(content, heading_style="ATX")
    else:
        return clean_markdown_text(content)

def format_sources_md(source_info):
    if not source_info:
        return ""
    sources_text = "\n\n## Sources\n"
    for idx, source in enumerate(source_info, start=1):
        title = source.get("title", "Inconnu")
        url = source.get("url", "#")
        date = source.get("date")
        if date:
            sources_text += f"[{idx}] [{title} ({date})]({url})\n"
        else:
            sources_text += f"[{idx}] [{title}]({url})\n"
    return sources_text

def format_sources_html(source_info):
    if not source_info:
        return ""
    html_str = "<h2>Sources</h2><ul>"
    for idx, source in enumerate(source_info, start=1):
        title = html.escape(source.get("title", "Inconnu"))
        url = html.escape(source.get("url", "#"))
        date = source.get("date")
        date_str = f" ({date})" if date else ""
        html_str += f'''
        <li>
            <span class="ref-number">[{idx}]</span>
            <a href="{url}">{title}{date_str}</a>
        </li>'''
    html_str += "</ul>"
    return html_str

def wrap_html(html_body: str, title: str, css_content: str, source_info=None, for_pdf=False) -> str:
    # Protection contre les valeurs None
    html_body = html_body or ""
    title = title or "Document Sans Titre"
    css_content = css_content or ""
    
    sources_html = format_sources_html(source_info)
    num_sources = len(source_info) if source_info else 0
    current_date = datetime.now().strftime("%d %B %Y")
    short_title = title[:50] + "..." if len(title) > 50 else title
    
    if for_pdf:
        # Charger le logo en base64
        logo_base64 = load_logo_base64()
        if logo_base64:
            logo_html = f'<img src="data:image/png;base64,{logo_base64}" class="cover-logo-img" />'
        else:
            logo_html = '<div class="cover-brand-text">SynthetIA</div>'
        
        # Premium Cover Page - Titre en haut, logo en bas
        cover_page = f'''
        <div class="cover-page">
            <div class="cover-title">{html.escape(title)}</div>
            <div class="cover-subtitle">Synthèse Automatisée</div>
            
            <div class="cover-meta">
                <div class="cover-meta-item">
                    <span class="cover-meta-label">Date de génération :</span> {current_date}
                </div>
                <div class="cover-meta-item">
                    <span class="cover-meta-label">Sources analysées :</span> {num_sources} vidéo(s)
                </div>
            </div>
            
            <div class="cover-footer">
                {logo_html}
            </div>
        </div>
        '''
        
        header_div = f'''
        <div id="headerContent">
            {html.escape(short_title)}
        </div>
        '''
        
        footer_div = f'''
        <div id="footerContent">
            <span class="footer-brand">SynthetIA</span> • {current_date} • Page <pdf:pagenumber /> / <pdf:pagecount />
        </div>
        '''
        
        toc_div = '''
        <div id="toc-container">
            <h1>Sommaire</h1>
            <pdf:toc />
        </div>
        '''
    else:
        cover_page = f'<div class="doc-title">{html.escape(title)}</div>'
        header_div = ""
        footer_div = ""
        toc_div = ""
        
    return f"""
    <html>
    <head>
        <meta charset="utf-8">
        <style>
        {css_content}
        </style>
    </head>
    <body>
        {cover_page}
        
        {toc_div}
        
        <div class="content">
            {html_body}
        </div>
        
        <div id="sources">
            {sources_html}
        </div>
        
        {header_div}
        {footer_div}
    </body>
    </html>
    """

# --- Strategies ---

class ExportStrategy(ABC):
    @abstractmethod
    def export(self, summary: str, output_file: str, title: str, source_info=None):
        pass

class MarkdownStrategy(ExportStrategy):
    def export(self, summary: str, output_file: str, title: str, source_info=None):
        title = title or "Document Sans Titre"
        summary = summary or ""
        content = f"# {title}\n\n"
        content += to_markdown(summary)
        content += format_sources_md(source_info)
        with open(output_file, "w", encoding="utf-8") as f:
            f.write(content)

class TextStrategy(ExportStrategy):
    def export(self, summary: str, output_file: str, title: str, source_info=None):
        title = title or "Document Sans Titre"
        summary = summary or ""
        content = f"{title}\n{'='*len(title)}\n\n"
        content += to_markdown(summary)
        content += format_sources_md(source_info)
        with open(output_file, "w", encoding="utf-8") as f:
            f.write(content)

class HTMLStrategy(ExportStrategy):
    def __init__(self, css_content):
        self.css_content = css_content

    def export(self, summary: str, output_file: str, title: str, source_info=None):
        html_body = to_html(summary)
        full_html = wrap_html(html_body, title, self.css_content, source_info, for_pdf=False)
        with open(output_file, "w", encoding="utf-8") as f:
            f.write(full_html)

class PDFStrategy(ExportStrategy):
    def __init__(self, css_content):
        self.css_content = css_content

    def export(self, summary: str, output_file: str, title: str, source_info=None):
        html_body = to_html(summary)
        full_html = wrap_html(html_body, title, self.css_content, source_info, for_pdf=True)
        
        pdf_file = io.BytesIO()
        pisa.CreatePDF(full_html, dest=pdf_file)
        
        with open(output_file, "wb") as f:
            f.write(pdf_file.getvalue())


class DocxStrategy(ExportStrategy):
    """Export vers Word (.docx) via python-docx."""
    
    def export(self, summary: str, output_file: str, title: str, source_info=None):
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.style import WD_STYLE_TYPE
        except ImportError:
            raise ImportError("python-docx requis. Installez avec: pip install python-docx")
        
        doc = Document()
        
        # Titre principal
        title_para = doc.add_heading(title or "Document Sans Titre", level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Date de génération
        date_para = doc.add_paragraph(f"Généré le {datetime.now().strftime('%d %B %Y')}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()  # Espace
        
        # Contenu - convertir en Markdown puis traiter
        md_content = to_markdown(summary)
        
        # Parser le Markdown ligne par ligne
        lines = md_content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # Titres H1-H4
            if line.startswith('#### '):
                doc.add_heading(line[5:], level=4)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            # Listes
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line[0].isdigit() and '. ' in line[:4]:
                doc.add_paragraph(line.split('. ', 1)[1], style='List Number')
            # Paragraphe normal
            else:
                # Nettoyer le gras/italique basique
                text = line.replace('**', '').replace('__', '').replace('*', '').replace('_', '')
                doc.add_paragraph(text)
        
        # Sources
        if source_info:
            doc.add_heading("Sources", level=1)
            for idx, source in enumerate(source_info, start=1):
                title_src = source.get("title", "Inconnu")
                url = source.get("url", "")
                date = source.get("date", "")
                date_str = f" ({date})" if date else ""
                doc.add_paragraph(f"[{idx}] {title_src}{date_str}", style='List Bullet')
                if url:
                    doc.add_paragraph(f"    {url}")
        
        doc.save(output_file)


class PptxStrategy(ExportStrategy):
    """Export vers PowerPoint (.pptx) via python-pptx."""
    
    def export(self, summary: str, output_file: str, title: str, source_info=None):
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN
        except ImportError:
            raise ImportError("python-pptx requis. Installez avec: pip install python-pptx")
        
        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9
        prs.slide_height = Inches(7.5)
        
        # Slide de titre
        title_slide_layout = prs.slide_layouts[6]  # Blank
        slide = prs.slides.add_slide(title_slide_layout)
        
        # Titre centré
        left = Inches(0.5)
        top = Inches(2.5)
        width = Inches(12.333)
        height = Inches(1.5)
        title_box = slide.shapes.add_textbox(left, top, width, height)
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = title or "Synthèse"
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 75, 75)  # Orange-rouge SynthetIA
        p.alignment = PP_ALIGN.CENTER
        
        # Sous-titre
        sub_box = slide.shapes.add_textbox(left, Inches(4), width, Inches(0.5))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"Généré par SynthetIA • {datetime.now().strftime('%d %B %Y')}"
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(150, 150, 150)
        p.alignment = PP_ALIGN.CENTER
        
        # Parser le contenu
        md_content = to_markdown(summary)
        sections = self._parse_sections(md_content)
        
        # Créer une slide par section H2
        for section_title, section_content in sections:
            self._add_content_slide(prs, section_title, section_content)
        
        # Slide Sources
        if source_info:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            # Titre
            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = "Sources"
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 75, 75)
            
            # Liste sources
            content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5.5))
            tf = content_box.text_frame
            tf.word_wrap = True
            for idx, source in enumerate(source_info, start=1):
                p = tf.add_paragraph() if idx > 1 else tf.paragraphs[0]
                p.text = f"[{idx}] {source.get('title', 'Inconnu')}"
                p.font.size = Pt(14)
                p.space_after = Pt(6)
        
        prs.save(output_file)
    
    def _parse_sections(self, md_content: str):
        """Parse le Markdown en sections (titre H2 + contenu)."""
        sections = []
        current_title = "Points Clés"
        current_content = []
        
        for line in md_content.split('\n'):
            if line.startswith('## '):
                if current_content:
                    sections.append((current_title, '\n'.join(current_content)))
                current_title = line[3:].strip()
                current_content = []
            elif line.startswith('# '):
                # H1 devient aussi une section
                if current_content:
                    sections.append((current_title, '\n'.join(current_content)))
                current_title = line[2:].strip()
                current_content = []
            else:
                current_content.append(line)
        
        if current_content:
            sections.append((current_title, '\n'.join(current_content)))
        
        return sections[:10]  # Max 10 slides
    
    def _add_content_slide(self, prs, section_title: str, content: str):
        """Ajoute une slide avec titre et contenu."""
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Titre
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = section_title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 75, 75)
        
        # Contenu
        content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(12), Inches(5.8))
        tf = content_box.text_frame
        tf.word_wrap = True
        
        lines = [l.strip() for l in content.split('\n') if l.strip()]
        first = True
        for line in lines[:12]:  # Max 12 lignes par slide
            # Nettoyer le formatage Markdown
            clean = line.replace('**', '').replace('*', '').replace('###', '').replace('##', '').replace('#', '')
            if clean.startswith('- ') or clean.startswith('* '):
                clean = '• ' + clean[2:]
            
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = clean
            p.font.size = Pt(16)
            p.space_after = Pt(8)

# --- Context Class ---

class Exporter:
    def __init__(self, output_dir: str):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)
        self.css_content = load_css()
        
        # Initialize strategies
        self.strategies = {
            "md": MarkdownStrategy(),
            "txt": TextStrategy(),
            "html": HTMLStrategy(self.css_content),
            "pdf": PDFStrategy(self.css_content),
            "docx": DocxStrategy(),
            "pptx": PptxStrategy()
        }

    def save_summary(self, summary: str, title: str, fmt: str, source_info=None):
        if fmt not in self.strategies:
            raise ValueError(f"Format non supporté : {fmt}. Utilisez: {', '.join(self.strategies.keys())}")
            
        slug = slugify(title)
        date_str = datetime.now().strftime("%Y-%m-%d")
        filename = f"{slug}_{date_str}.{fmt}"
        output_file = os.path.join(self.output_dir, filename)

        self.strategies[fmt].export(summary, output_file, title, source_info)

        return output_file

    def generate_pdf_bytes(self, summary: str, title: str, source_info=None) -> bytes:
        # Protection contre les valeurs None
        summary = summary or ""
        title = title or "Document Sans Titre"
        
        html_body = to_html(summary)
        full_html = wrap_html(html_body, title, self.css_content, source_info, for_pdf=True)
        
        pdf_file = io.BytesIO()
        pisa.CreatePDF(full_html, dest=pdf_file)
        return pdf_file.getvalue()
